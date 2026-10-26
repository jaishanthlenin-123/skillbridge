import re
import fitz # PyMuPDF
import docx
from pathlib import Path
from django.conf import settings
from .models import Resume, ResumeSection, ResumeSkill, Skill, SectionType, ParsingStatus, ExtractionMethod

SECTION_HEADER_PATTERNS = [
    (SectionType.SUMMARY, re.compile(r'^(summary|professional summary|profile|about me|objective)\b', re.IGNORECASE)),
    (SectionType.EXPERIENCE, re.compile(r'^(experience|work experience|employment history|work history|professional experience)\b', re.IGNORECASE)),
    (SectionType.EDUCATION, re.compile(r'^(education|academic background|academics|qualifications)\b', re.IGNORECASE)),
    (SectionType.SKILLS, re.compile(r'^(skills|technical skills|technologies|core competencies|expertise|tools)\b', re.IGNORECASE)),
    (SectionType.PROJECTS, re.compile(r'^(projects|personal projects|technical projects|key projects)\b', re.IGNORECASE)),
    (SectionType.CERTIFICATIONS, re.compile(r'^(certifications|certificates|licenses|awards|honors)\b', re.IGNORECASE)),
]

def extract_raw_text_from_file(file_path: str, ext: str) -> str:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    text = ""
    if ext == 'pdf':
        doc = fitz.open(file_path)
        for page in doc:
            text += page.get_text() + "\n"
        doc.close()
    elif ext == 'docx':
        doc = docx.Document(file_path)
        for para in doc.paragraphs:
            text += para.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + " "
                text += "\n"
    else:
        raise ValueError(f"Unsupported format: {ext}")

    # Normalize whitespace & artifacts
    text = re.sub(r'\r\n|\r', '\n', text)
    text = re.sub(r'[ \t]+', ' ', text)
    text = re.sub(r'\n{3,}', '\n\n', text).strip()
    return text

def segment_sections(text: str) -> list[dict]:
    lines = [line.strip() for line in text.split('\n') if line.strip()]
    if not lines:
        return []

    sections = []
    current_type = SectionType.SUMMARY
    current_lines = []

    for line in lines:
        matched_type = None
        # Check if line looks like a header (short line, matches regex)
        if len(line) < 60:
            for sec_type, pattern in SECTION_HEADER_PATTERNS:
                if pattern.search(line):
                    matched_type = sec_type
                    break

        if matched_type:
            if current_lines:
                sections.append({
                    'type': current_type,
                    'content': '\n'.join(current_lines).strip()
                })
                current_lines = []
            current_type = matched_type
        else:
            current_lines.append(line)

    if current_lines:
        sections.append({
            'type': current_type,
            'content': '\n'.join(current_lines).strip()
        })

    return sections

def process_resume_document(resume: Resume) -> bool:
    """
    Synchronous document extraction and section parsing.
    """
    resume.parsing_status = ParsingStatus.PROCESSING
    resume.save(update_fields=['parsing_status'])

    try:
        abs_path = Path(settings.PRIVATE_STORAGE_ROOT) / resume.private_storage_key
        ext = resume.original_filename.split('.')[-1].lower()

        raw_text = extract_raw_text_from_file(str(abs_path), ext)

        if not raw_text or len(raw_text.strip()) < 20:
            resume.parsing_status = ParsingStatus.FAILED
            resume.error_details = (
                "Scanned or image-only PDF detected with no extractable text. "
                "Please upload a document with searchable text."
            )
            resume.save()
            return False

        resume.raw_text = raw_text

        # Purge existing sections/skills if re-processing
        resume.sections.all().delete()
        resume.skills.all().delete()

        # Parse sections
        segmented = segment_sections(raw_text)
        for idx, sec in enumerate(segmented):
            ResumeSection.objects.create(
                resume=resume,
                section_type=sec['type'],
                content=sec['content'],
                order_index=idx
            )

        # Match skills from canonical Skill and SkillAlias database
        from .seed_taxonomy import extract_skills_from_text
        found_skills = extract_skills_from_text(raw_text)

        for item in found_skills:
            canonical_skill = Skill.objects.filter(name__iexact=item['name']).first()
            ResumeSkill.objects.update_or_create(
                resume=resume,
                skill_name=item['name'],
                defaults={
                    'skill': canonical_skill,
                    'evidence_text': item.get('evidence', '')[:500],
                    'extraction_method': ExtractionMethod.TAXONOMY_PHRASE,
                    'confidence': item.get('confidence', 0.95),
                }
            )

        resume.parsing_status = ParsingStatus.COMPLETED
        resume.error_details = ''
        resume.save()
        return True

    except Exception as e:
        resume.parsing_status = ParsingStatus.FAILED
        resume.error_details = f"Parsing error: {str(e)}"
        resume.save()
        return False
